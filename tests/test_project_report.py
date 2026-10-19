"""Methods text, references and the combined project report.

Run from backend/:  python -m unittest discover -s tests -t .
"""
import io
import re
import unittest

from tests.support import client_and_headers  # must come before other app imports

from docx import Document  # noqa: E402

from app import dataprep, methods_text  # noqa: E402

SAVED_PARAMETERS = {
    "descriptive": {"variables": ["Yield", "Height"], "group_var": "Variety"},
    "correlation": {"method": "spearman", "variables": ["Yield", "Height", "Protein"]},
    "parametric": {"test": "independent_t", "variables": ["Yield_A", "Yield_B"]},
    "nonparametric": {"test": "kruskal_wallis", "variables": ["T1", "T2", "T3"]},
    "anova": {"design": "splitplot", "dependent": "Yield", "factor_a": "Irrigation", "factor_b": "Nitrogen",
              "block": "Rep", "posthoc": "duncan"},
    "regression": {"type": "plsr", "dependent": ["Yield"], "independent": ["N", "P", "K"], "components": 2},
    "pca": {"variables": ["pH", "OC", "EC"], "scale": True, "varimax": 2},
    "clustering": {"method": "hierarchical", "k": 3, "variables": ["pH", "OC"], "linkage": "ward", "distance": "euclidean"},
    "trend": {"variables": ["Rainfall"], "time": "Year", "season": "Month", "alpha": 0.05},
    "gxe": {"genotype": "Genotype", "environment": "Location", "response": "Yield", "replicate": "Rep"},
    "divergence": {"genotype": "Genotype", "traits": ["Plant height", "Grain yield"]},
    "timeseries": {"value": "Price", "time": "Month", "season_length": 12, "horizon": 6},
    "sem": {"type": "cb", "latents": [{"name": "Soil health"}, {"name": "Productivity"}]},
    "imputation": {"columns": ["Yield"], "predictors": ["Tillers", "Irrigation"], "m": 20, "max_iter": 10, "seed": 2026,
                   "analyses": {"correlation": {"variables": ["Yield", "Tillers"], "method": "spearman"},
                                "anova": {"response": "Yield", "factors": ["Irrigation"], "block": "Rep"}}},
}


class MethodsTextTests(unittest.TestCase):
    def test_every_module_gets_a_specific_paragraph_with_known_references(self):
        for module, parameters in SAVED_PARAMETERS.items():
            described = methods_text.describe_analysis(module, parameters, {})
            paragraph = described["paragraph"]
            self.assertGreater(len(paragraph), 80, module)
            self.assertNotIn("None", paragraph, module)
            for key in described["references"]:
                self.assertIn(key, methods_text.REFERENCES, module)

    def test_paragraphs_follow_the_settings(self):
        welch = methods_text.describe_analysis("parametric", SAVED_PARAMETERS["parametric"], {"test_used": "Welch's t-Test"})
        self.assertIn("Welch's t-test", welch["paragraph"])
        self.assertIn("welch1947", welch["references"])
        student = methods_text.describe_analysis("parametric", SAVED_PARAMETERS["parametric"], {"test_used": "Student's t-Test"})
        self.assertIn("Student's", student["paragraph"])

        anova = methods_text.describe_analysis("anova", SAVED_PARAMETERS["anova"])
        self.assertIn("split-plot", anova["paragraph"])
        self.assertIn("Irrigation and Nitrogen", anova["paragraph"])
        self.assertIn("Duncan's multiple range test", anova["paragraph"])
        self.assertIn("duncan1955", anova["references"])
        self.assertNotIn("tukey1949", anova["references"])

        spearman = methods_text.describe_analysis("correlation", SAVED_PARAMETERS["correlation"])
        self.assertIn("Spearman", spearman["paragraph"])
        pearson = methods_text.describe_analysis("correlation", {"method": "pearson", "variables": ["a", "b"]})
        self.assertNotIn("spearman1904", pearson["references"])

        seasonal = methods_text.describe_analysis("trend", SAVED_PARAMETERS["trend"])
        self.assertIn("seasonal Mann-Kendall", seasonal["paragraph"])
        self.assertIn("hirsch1982", seasonal["references"])
        plain = methods_text.describe_analysis("trend", {"variables": ["Rainfall"]})
        self.assertNotIn("hirsch1982", plain["references"])

        mi = methods_text.describe_analysis("imputation", SAVED_PARAMETERS["imputation"])
        self.assertIn("m = 20", mi["paragraph"])
        self.assertIn("D1", mi["paragraph"])
        self.assertIn("Fisher's z", mi["paragraph"])
        for key in ("rubin1987", "barnard1999", "li1991", "fisher1915", "fieller1957"):
            self.assertIn(key, mi["references"])

        pls = methods_text.describe_analysis("sem", {"type": "pls", "latents": [{"name": "A"}, {"name": "B"}]})
        self.assertIn("bootstrapping", pls["paragraph"])
        self.assertIn("henseler2015", pls["references"])

    def test_unknown_module_does_not_fail(self):
        self.assertEqual(methods_text.describe_analysis("teleport")["references"], [])

    def test_every_recipe_step_can_be_described(self):
        examples = {
            "drop_missing": {"axis": "columns", "max_missing_fraction": 0.4},
            "fill_missing": {"columns": ["Yield"], "method": "median", "group_by": "Treatment"},
            "impute_model": {"columns": ["Yield"], "method": "knn", "n_neighbors": 4},
            "filter_rows": {"column": "Height", "op": "between", "value": 80, "value2": 120},
            "handle_outliers": {"columns": ["Yield"], "method": "mad", "action": "winsorize"},
            "compute_column": {"name": "Yield_ha", "formula": "Yield * 500"},
            "bin_numeric": {"column": "pH", "method": "custom", "edges": [0, 6.5, 7.5, 14]},
            "rename_columns": {"mapping": {"Ht": "Height"}},
        }
        for kind in dataprep.STEPS:
            sentence, keys = methods_text.describe_step({"type": kind, **examples.get(kind, {})})
            self.assertTrue(sentence.endswith("."), kind)
            self.assertNotIn(" None", sentence, kind)
            for key in keys:
                self.assertIn(key, methods_text.REFERENCES)
        self.assertIn("40%", methods_text.describe_step({"type": "drop_missing", "axis": "columns", "max_missing_fraction": 0.4})[0])
        sentence, keys = methods_text.describe_step(examples["handle_outliers"] | {"type": "handle_outliers"})
        self.assertIn("median absolute deviation", sentence)
        self.assertEqual(keys, ["iglewicz1993"])

    def test_reference_list_is_unique_sorted_and_complete(self):
        refs = methods_text.reference_list(["tukey1949", "duncan1955", "tukey1949", "not-a-key"])
        self.assertEqual(len(refs), 2)
        self.assertEqual(refs, sorted(refs, key=str.lower))
        for key, text in methods_text.REFERENCES.items():
            self.assertRegex(text, r"\(\d{4}\)", key)  # every entry has a year


def docx_text(content):
    document = Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts), document


class ProjectReportApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client, cls.headers = client_and_headers("report-owner@example.com")
        cls.other = client_and_headers("report-other@example.com")[1]

    def api(self, method, path, headers=None, **kwargs):
        return self.client.request(method, f"/workspace{path}", headers=headers or self.headers, **kwargs)

    def make_project(self):
        project = self.api("POST", "/projects", json={"name": "Wheat N trial"}).json()
        csv = b"Rep,Irrigation,Nitrogen,Yield\n1,I1,N0,3.1\n1,I1,N60,4.0\n"
        self.api("POST", f"/projects/{project['id']}/datasets", files={"file": ("wheat.csv", csv, "text/csv")})
        stored_report = {
            "title": "ANOVA report",
            "sections": [
                {"heading": "ANOVA table", "tables": [{"columns": ["Source", "df", "F", "p"], "rows": [["Nitrogen", 2, 18.4, 0.0004]]}]},
                {"heading": "Mean comparison", "paragraphs": ["N120 gave the highest mean yield."]},
            ],
            "notes": ["Split-plot error terms were used."],
        }
        anova = self.api("POST", f"/projects/{project['id']}/results", json={
            "module": "anova", "title": "Yield ANOVA", "parameters": SAVED_PARAMETERS["anova"],
            "result": {"posthoc_method": "duncan"}, "report": stored_report}).json()
        older = self.api("POST", f"/projects/{project['id']}/results", json={
            "module": "correlation", "title": "Trait correlations", "parameters": SAVED_PARAMETERS["correlation"],
            "result": {"matrix": {}}}).json()
        recipe = self.api("POST", f"/projects/{project['id']}/recipes", json={
            "name": "Kharif cleaning", "source_dataset": "wheat.csv",
            "steps": [{"type": "drop_duplicates"},
                      {"type": "handle_outliers", "columns": ["Yield"], "method": "iqr", "action": "to_missing"}]}).json()
        return project, anova, older, recipe

    def test_report_contains_every_part(self):
        project, anova, older, recipe = self.make_project()
        response = self.api("POST", f"/projects/{project['id']}/report", json={})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("wordprocessingml", response.headers["content-type"])
        self.assertIn("Wheat_N_trial", response.headers["content-disposition"])
        text, document = docx_text(response.content)

        headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
        for expected in ("Data", "Data preparation", "Statistical methods", "Result 1: Yield ANOVA",
                         "ANOVA table", "Mean comparison", "Result 2: Trait correlations", "Settings", "Activity log", "References"):
            self.assertIn(expected, headings)
        # the stored section is a sub-section of its result
        levels = {p.text: p.style.name for p in document.paragraphs if p.style.name.startswith("Heading")}
        self.assertEqual(levels["Result 1: Yield ANOVA"], "Heading 2")
        self.assertEqual(levels["ANOVA table"], "Heading 3")

        self.assertIn("wheat.csv", text)
        self.assertIn("Recipe \"Kharif cleaning\"", text)
        self.assertIn("Tukey's fences", text)
        self.assertIn("Duncan's multiple range test", text)
        self.assertIn("Nitrogen | 2 | 18.4000 | 4.000e-04", text)  # the report's usual number formatting
        self.assertIn("saved before project reports kept them", text)
        self.assertIn("Duncan, D. B. (1955)", text)
        self.assertIn("Tukey, J. W. (1977)", text)
        self.assertIn("Spearman, C. (1904)", text)
        self.assertIn("Dataset uploaded", text)
        self.assertIn("Recipe saved", text)

    def test_choosing_what_goes_in(self):
        project, anova, older, recipe = self.make_project()
        response = self.api("POST", f"/projects/{project['id']}/report", json={
            "result_ids": [anova["id"]], "recipe_ids": [], "include_activity": False, "include_references": False,
            "title": "Chapter 4 tables"})
        text, document = docx_text(response.content)
        headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
        self.assertIn("Chapter 4 tables", headings)
        self.assertNotIn("Result 2: Trait correlations", headings)
        self.assertNotIn("Data preparation", headings)
        self.assertNotIn("Activity log", headings)
        self.assertNotIn("References", headings)

    def test_methods_endpoint(self):
        project, anova, older, recipe = self.make_project()
        body = self.api("GET", f"/projects/{project['id']}/methods").json()
        self.assertEqual([a["title"] for a in body["analyses"]], ["Yield ANOVA", "Trait correlations"])
        self.assertEqual(body["preparation"][0]["name"], "Kharif cleaning")
        self.assertEqual(len(body["preparation"][0]["sentences"]), 2)
        self.assertIn("SciPy", body["software"])
        self.assertTrue(any(ref.startswith("Duncan") for ref in body["references"]))
        self.assertTrue(any(ref.startswith("Virtanen") for ref in body["references"]))

    def test_report_content_counts_toward_storage_and_is_private(self):
        project = self.api("POST", "/projects", json={"name": "Size check"}).json()
        plain = self.api("POST", f"/projects/{project['id']}/results", json={
            "module": "pca", "title": "a", "parameters": {}, "result": {"x": 1}}).json()
        with_report = self.api("POST", f"/projects/{project['id']}/results", json={
            "module": "pca", "title": "b", "parameters": {}, "result": {"x": 1},
            "report": {"sections": [{"heading": "Eigenvalues", "paragraphs": ["x" * 2000]}]}}).json()
        self.assertGreater(with_report["size_bytes"], plain["size_bytes"] + 2000)
        self.assertEqual(self.api("POST", f"/projects/{project['id']}/report", headers=self.other, json={}).status_code, 404)
        self.assertEqual(self.api("GET", f"/projects/{project['id']}/methods", headers=self.other).status_code, 404)

    def test_empty_project_still_produces_a_document(self):
        project = self.api("POST", "/projects", json={"name": "Empty"}).json()
        response = self.api("POST", f"/projects/{project['id']}/report", json={})
        self.assertEqual(response.status_code, 200)
        self.assertIn("no datasets, recipes or saved analyses", docx_text(response.content)[0])

    def test_a_malformed_stored_report_falls_back_to_settings(self):
        project = self.api("POST", "/projects", json={"name": "Broken"}).json()
        self.api("POST", f"/projects/{project['id']}/results", json={
            "module": "anova", "title": "Odd", "parameters": {"design": "oneway"}, "result": {},
            "report": {"sections": [{"no_heading": True}]}})
        response = self.api("POST", f"/projects/{project['id']}/report", json={})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(re.search(r"Settings", docx_text(response.content)[0]))


if __name__ == "__main__":
    unittest.main()
