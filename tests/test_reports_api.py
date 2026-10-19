"""Server-side Word report generation (/reports/docx).

Run from backend/:  python -m unittest discover -s tests -t .
"""
import base64
import io
import unittest

from tests.support import client_and_headers  # must come before other app imports

from docx import Document  # noqa: E402
from PIL import Image  # noqa: E402

from app.reports import format_cell  # noqa: E402


def png_base64():
    buffer = io.BytesIO()
    Image.new("RGB", (120, 80), (79, 70, 229)).save(buffer, format="PNG")
    return base64.b64encode(buffer.getvalue()).decode()


def sample_report(**overrides):
    report = {
        "title": "One-Way ANOVA Report",
        "test_applied": "CRD one-way ANOVA with Tukey's HSD",
        "dataset_name": "wheat_yield.csv",
        "filename": "Wheat yield / ANOVA",
        "sections": [
            {"heading": "ANOVA Table", "paragraphs": ["Treatments differ significantly."],
             "tables": [{"columns": ["Source", "df", "SS", "p-value"],
                         "rows": [["Treatment", 3, 152.4321, 0.00012], ["Error", 16, 40.5, None]],
                         "caption": "Table 1. Analysis of variance", "note": "p-values are two-sided."}]},
            {"heading": "Mean Comparison", "bullets": ["T4 has the highest mean", "T1 and T2 do not differ"],
             "images": [{"data": f"data:image/png;base64,{png_base64()}", "caption": "Figure 1. Treatment means"}]},
        ],
        "notes": ["Levene's test p = 0.41"],
    }
    report.update(overrides)
    return report


class ReportApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client, cls.headers = client_and_headers()

    def post(self, body, headers=None):
        return self.client.post("/reports/docx", json=body, headers=self.headers if headers is None else headers)

    def test_generates_a_real_word_document(self):
        response = self.post(sample_report())
        self.assertEqual(response.status_code, 200, response.text[:300])
        self.assertTrue(response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
        self.assertIn('filename="Wheat_yield_ANOVA.docx"', response.headers["content-disposition"])

        doc = Document(io.BytesIO(response.content))
        text = "\n".join(p.text for p in doc.paragraphs)
        for expected in ("One-Way ANOVA Report", "ANOVA Table", "Treatments differ significantly.",
                         "T4 has the highest mean", "Figure 1. Treatment means", "Notes", "Levene's test p = 0.41"):
            self.assertIn(expected, text)

        meta, anova = doc.tables[0], doc.tables[1]
        self.assertIn("wheat_yield.csv", [c.text for row in meta.rows for c in row.cells])
        self.assertEqual([c.text for c in anova.rows[0].cells], ["Source", "df", "SS", "p-value"])
        self.assertEqual([c.text for c in anova.rows[1].cells], ["Treatment", "3", "152.4321", "1.200e-04"])
        self.assertEqual(anova.rows[2].cells[3].text, "–")
        self.assertEqual(len(doc.inline_shapes), 2)  # logo + chart

    def test_number_formatting(self):
        self.assertEqual(format_cell(0.0), "0.0000")
        self.assertEqual(format_cell(12.5), "12.5000")
        self.assertEqual(format_cell(2.5e-6), "2.500e-06")
        self.assertEqual(format_cell(float("nan")), "–")
        self.assertEqual(format_cell(True), "Yes")
        self.assertEqual(format_cell("p < 0.001"), "p < 0.001")

    def test_rejects_ragged_tables_and_bad_images(self):
        ragged = sample_report(sections=[{"heading": "Bad", "tables": [{"columns": ["a", "b"], "rows": [[1]]}]}])
        self.assertEqual(self.post(ragged).status_code, 400)
        bad_image = sample_report(sections=[{"heading": "Chart", "images": [{"data": "not-base64!!"}]}])
        self.assertEqual(self.post(bad_image).status_code, 400)
        not_png = sample_report(sections=[{"heading": "Chart", "images": [{"data": base64.b64encode(b"GIF89a....").decode()}]}])
        self.assertEqual(self.post(not_png).status_code, 400)

    def test_requires_sign_in(self):
        self.assertEqual(self.post(sample_report(), headers={}).status_code, 401)


if __name__ == "__main__":
    unittest.main()
