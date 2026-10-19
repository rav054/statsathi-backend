"""Word (.docx) reports generated on the server from structured report content.

Each analysis screen sends the same simple structure (title, sections with paragraphs,
bullet points, tables and chart images), so every module gets consistently formatted,
genuine Word documents from one implementation.
"""
import base64
import binascii
import io
import re
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Union

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from .auth import get_current_user
from .models import User

router = APIRouter(prefix="/reports", tags=["Reports"])

LOGO_PATH = Path(__file__).resolve().parent / "assets" / "logo.png"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
MAX_IMAGE_BYTES = 8 * 1024 * 1024
MAX_TABLE_CELLS = 50_000

Cell = Union[str, int, float, bool, None]


class ReportTable(BaseModel):
    columns: List[str] = Field(min_length=1, max_length=60)
    rows: List[List[Cell]] = Field(default_factory=list, max_length=5000)
    caption: Optional[str] = Field(None, max_length=300)
    note: Optional[str] = Field(None, max_length=1000)


class ReportImage(BaseModel):
    data: str  # base64 PNG or JPEG, optionally as a data: URI
    caption: Optional[str] = Field(None, max_length=300)
    width_inches: float = Field(6.0, gt=1.0, le=6.5)


class ReportSection(BaseModel):
    heading: str = Field(min_length=1, max_length=200)
    level: int = Field(2, ge=1, le=3)  # 2 for a section, 3 for a sub-section
    paragraphs: List[str] = []
    bullets: List[str] = []
    tables: List[ReportTable] = []
    images: List[ReportImage] = []


class ReportIn(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    test_applied: Optional[str] = Field(None, max_length=300)
    dataset_name: Optional[str] = Field(None, max_length=255)
    filename: str = Field("StatSathi_Report", max_length=150)
    sections: List[ReportSection] = Field(min_length=1, max_length=400)
    notes: List[str] = []


def format_cell(value: Cell) -> str:
    """Numbers get fixed decimals (scientific notation when very small or large); text is kept as sent."""
    if value is None:
        return "–"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value != value:  # NaN
            return "–"
        if value == 0 or 1e-3 <= abs(value) < 1e6:
            return f"{value:.4f}"
        return f"{value:.3e}"
    return str(value)


def _decode_image(image: ReportImage) -> bytes:
    data = image.data.split(",", 1)[1] if image.data.startswith("data:") else image.data
    try:
        raw = base64.b64decode(data, validate=True)
    except (binascii.Error, ValueError):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="A report image is not valid base64 data.")
    if len(raw) > MAX_IMAGE_BYTES:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="A report image is larger than 8 MB.")
    if not (raw.startswith(b"\x89PNG\r\n\x1a\n") or raw.startswith(b"\xff\xd8\xff")):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Report images must be PNG or JPEG.")
    return raw


def _validate(report: ReportIn):
    cells = 0
    for section in report.sections:
        for table in section.tables:
            for row in table.rows:
                if len(row) != len(table.columns):
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail=f"Table in section '{section.heading}' has a row with {len(row)} cells but {len(table.columns)} columns."
                    )
            cells += len(table.rows) * len(table.columns)
    if cells > MAX_TABLE_CELLS:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="The report tables are too large.")


def _shade(cell, hex_fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_fill)
    properties.append(shading)


def _repeat_header(row):
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _add_field(paragraph, instruction):
    """Insert a Word field (e.g. PAGE) so Word fills in the value."""
    run = paragraph.add_run()
    for tag, text in (("w:fldChar", "begin"), ("w:instrText", instruction), ("w:fldChar", "end")):
        element = OxmlElement(tag)
        if tag == "w:fldChar":
            element.set(qn("w:fldCharType"), text)
        else:
            element.set(qn("xml:space"), "preserve")
            element.text = instruction
        run._r.append(element)


def _styled_run(paragraph, text, size=None, bold=False, italic=False, color=None):
    run = paragraph.add_run(text)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return run


def _add_table(doc, table: ReportTable):
    if table.caption:
        caption = doc.add_paragraph()
        _styled_run(caption, table.caption, size=9.5, bold=True, color=SLATE)
    font_size = 8 if len(table.columns) > 7 else 9
    grid = doc.add_table(rows=1, cols=len(table.columns))
    grid.style = "Table Grid"
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER

    header = grid.rows[0]
    _repeat_header(header)
    for cell, name in zip(header.cells, table.columns):
        _shade(cell, "4F46E5")
        paragraph = cell.paragraphs[0]
        _styled_run(paragraph, name, size=font_size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    for row_index, values in enumerate(table.rows):
        cells = grid.add_row().cells
        for col_index, (cell, value) in enumerate(zip(cells, values)):
            if row_index % 2 == 1:
                _shade(cell, "F8FAFC")
            paragraph = cell.paragraphs[0]
            _styled_run(paragraph, format_cell(value), size=font_size, bold=(col_index == 0 and isinstance(value, str)))
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    if table.note:
        note = doc.add_paragraph()
        _styled_run(note, table.note, size=8.5, italic=True, color=MUTED)
    doc.add_paragraph()


def build_report(report: ReportIn, prepared_by: str) -> bytes:
    _validate(report)
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.8)
        section.left_margin = section.right_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    # Branded title block
    if LOGO_PATH.exists():
        logo = doc.add_paragraph()
        logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo.add_run().add_picture(str(LOGO_PATH), width=Inches(2.2))
    tagline = doc.add_paragraph()
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _styled_run(tagline, "StatSathi Research Suite · Your Research Analytics Companion", size=9, color=MUTED)

    title = doc.add_heading(report.title, level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = INDIGO

    meta_rows = [("Test applied", report.test_applied), ("Dataset", report.dataset_name),
                 ("Prepared by", prepared_by), ("Generated on", datetime.now().strftime("%d %B %Y, %H:%M"))]
    meta = doc.add_table(rows=0, cols=2)
    meta.style = "Table Grid"
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in meta_rows:
        if not value:
            continue
        cells = meta.add_row().cells
        _shade(cells[0], "F1F5F9")
        _styled_run(cells[0].paragraphs[0], label, size=9.5, bold=True, color=SLATE)
        _styled_run(cells[1].paragraphs[0], value, size=9.5)
    doc.add_paragraph()

    for section in report.sections:
        heading = doc.add_heading(section.heading, level=section.level)
        for run in heading.runs:
            run.font.color.rgb = INDIGO
        for text in section.paragraphs:
            doc.add_paragraph(text)
        for text in section.bullets:
            doc.add_paragraph(text, style="List Bullet")
        for table in section.tables:
            _add_table(doc, table)
        for image in section.images:
            picture = doc.add_paragraph()
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.add_run().add_picture(io.BytesIO(_decode_image(image)), width=Inches(image.width_inches))
            if image.caption:
                caption = doc.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _styled_run(caption, image.caption, size=9, italic=True, color=MUTED)

    if report.notes:
        heading = doc.add_heading("Notes", level=2)
        for run in heading.runs:
            run.font.color.rgb = INDIGO
        for text in report.notes:
            doc.add_paragraph(text, style="List Bullet")

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _styled_run(footer, "Generated by StatSathi · Curated by Ravi, PhD Scholar ICAR-IISS · Page ", size=8, color=MUTED)
    _add_field(footer, "PAGE")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@router.post("/docx")
def create_docx_report(report: ReportIn, current_user: User = Depends(get_current_user)):
    content = build_report(report, current_user.full_name)
    safe_name = re.sub(r"[^\w.-]+", "_", report.filename).strip("._") or "StatSathi_Report"
    return StreamingResponse(
        io.BytesIO(content),
        media_type=DOCX_MEDIA_TYPE,
        headers={
            "Content-Disposition": f'attachment; filename="{safe_name}.docx"',
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )
